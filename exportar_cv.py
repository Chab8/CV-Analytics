"""Exporta el CV ya limpio (ver crear_cv.limpiar_datos) a .docx y .txt.
Una sola columna y sin adornos: es lo que mejor leen los sistemas ATS."""
from io import BytesIO

from docx import Document
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ACENTO = RGBColor(0x0B, 0x4B, 0x64)
GRIS = RGBColor(0x5B, 0x6B, 0x73)


def _periodo(x):
    hasta = "Actualidad" if x.get("actual") else x.get("hasta", "")
    return " – ".join(v for v in (x.get("desde", ""), hasta) if v)


def _contacto(p):
    return [p[k] for k in ("email", "telefono", "ubicacion", "linkedin", "web") if p.get(k)]


def _bloque(titulo, sub="", fecha="", texto="", puntos=False):
    lineas = [l.strip() for l in texto.splitlines() if l.strip()]
    return {"titulo": titulo, "sub": sub, "fecha": fecha, "lineas": lineas, "puntos": puntos}


def _secciones(cv):
    """Estructura neutra [(titulo, [bloques])] que comparten el .docx y el .txt."""
    p = cv["personal"]
    out = []
    if p["resumen"]:
        out.append(("Perfil", [_bloque("", texto=p["resumen"])]))
    if cv["experiencia"]:
        out.append(("Experiencia laboral", [
            _bloque(e["puesto"], e["empresa"], _periodo(e), e["descripcion"], puntos=True)
            for e in cv["experiencia"]]))
    if cv["educacion"]:
        out.append(("Educación", [
            _bloque(e["titulo"], e["institucion"], _periodo(e), e["descripcion"])
            for e in cv["educacion"]]))
    if cv["habilidades"]:
        out.append(("Habilidades", [_bloque("", texto=", ".join(cv["habilidades"]))]))
    idiomas = [f'{i["idioma"]} ({i["nivel"]})' if i["nivel"] else i["idioma"]
               for i in cv["idiomas"] if i["idioma"]]
    if idiomas:
        out.append(("Idiomas", [_bloque("", texto=", ".join(idiomas))]))
    if cv["extras"]:
        out.append(("Cursos y certificaciones", [
            _bloque(x["titulo"], x["institucion"], x["anio"]) for x in cv["extras"]]))
    extra = cv.get("seccion_extra") or {}
    if extra.get("entradas"):
        out.append((extra.get("nombre") or "Otros", [
            _bloque(e["titulo"], "", e["fecha"], e["descripcion"]) for e in extra["entradas"]]))
    return out


# ---------- TXT ----------

def generar_txt(cv):
    p = cv["personal"]
    lineas = [p["nombre"]]
    if p["titulo"]:
        lineas.append(p["titulo"])
    contacto = _contacto(p)
    if contacto:
        lineas.append(" | ".join(contacto))
    for titulo, bloques in _secciones(cv):
        if lineas[-1] != "":
            lineas.append("")
        lineas += [titulo.upper(), "-" * len(titulo)]
        for b in bloques:
            cabecera = " ".join(filter(None, [b["titulo"], f"({b['fecha']})" if b["fecha"] else ""]))
            if cabecera:
                lineas.append(cabecera)
            if b["sub"]:
                lineas.append(b["sub"])
            lineas += [("- " if b["puntos"] else "") + l for l in b["lineas"]]
            lineas.append("")
    return ("\n".join(lineas).strip() + "\n").encode("utf-8")


# ---------- DOCX ----------

def _run(par, texto, size=None, color=None, bold=False, italic=False):
    r = par.add_run(texto)
    r.bold, r.italic = bold, italic
    if size:
        r.font.size = Pt(size)
    if color is not None:
        r.font.color.rgb = color
    return r


def _linea_inferior(par):
    borde = OxmlElement("w:pBdr")
    linea = OxmlElement("w:bottom")
    for k, v in (("val", "single"), ("sz", "6"), ("space", "1"), ("color", "0B4B64")):
        linea.set(qn("w:" + k), v)
    borde.append(linea)
    # pBdr tiene que ir antes de estos elementos para respetar el orden del esquema
    par._p.get_or_add_pPr().insert_element_before(
        borde,
        "w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap",
        "w:overflowPunct", "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN",
        "w:bidi", "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind",
        "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap", "w:jc",
        "w:textDirection", "w:textAlignment", "w:textboxTightWrap",
        "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange",
    )


def generar_docx(cv):
    d = Document()
    s = d.sections[0]
    s.left_margin = s.right_margin = Cm(2)
    s.top_margin = s.bottom_margin = Cm(1.8)

    normal = d.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.paragraph_format.space_after = Pt(2)

    p = cv["personal"]
    _run(d.add_paragraph(), p["nombre"], 22, ACENTO, bold=True)
    if p["titulo"]:
        _run(d.add_paragraph(), p["titulo"], 12, GRIS)
    contacto = _contacto(p)
    if contacto:
        _run(d.add_paragraph(), "  |  ".join(contacto), 9.5, GRIS)

    for titulo, bloques in _secciones(cv):
        h = d.add_paragraph()
        h.paragraph_format.space_before = Pt(12)
        h.paragraph_format.space_after = Pt(4)
        h.paragraph_format.keep_with_next = True
        _run(h, titulo, 12, ACENTO, bold=True)
        _linea_inferior(h)

        for b in bloques:
            if b["titulo"] or b["fecha"]:
                f = d.add_paragraph()
                f.paragraph_format.space_before = Pt(4)
                f.paragraph_format.keep_with_next = True
                f.paragraph_format.tab_stops.add_tab_stop(Cm(17), WD_TAB_ALIGNMENT.RIGHT)
                _run(f, b["titulo"], bold=True)
                if b["fecha"]:
                    _run(f, "\t" + b["fecha"], 9.5, GRIS)
            if b["sub"]:
                _run(d.add_paragraph(), b["sub"], color=ACENTO, italic=True)
            for l in b["lineas"]:
                par = d.add_paragraph(style="List Bullet") if b["puntos"] else d.add_paragraph()
                par.add_run(l)

    buf = BytesIO()
    d.save(buf)
    return buf.getvalue()